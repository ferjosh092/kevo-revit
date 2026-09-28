# -*- coding: utf-8 -*-
"""Libro de materiales e insumos del casco estructural - Vivienda multifamiliar San Jeronimo (Cusco).
Lee los resultados de solo lectura del IFC (q_*.json) y escribe un .xlsx con formulas vivas."""
import json, datetime, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment

OUT = "Lista_Materiales_Casco_SanJeronimo.xlsx"
Q = json.load(open("q_concreto.json"))
AE = json.load(open("q_acero_elem.json"))
CORTE = json.load(open("q_corte.json"))
AR = json.load(open("q_areas.json"))
GUID = json.load(open("q_guid.json"))
HOY = datetime.date(2026, 9, 28)

NIV = ["CIMENTACIÓN", "NIVEL 1", "NIVEL 2", "NIVEL 3", "NIVEL 4", "NIVEL 5"]
def nv(s): return "CIMENTACIÓN" if s.startswith("CIMENTACION") else s
ELEM = {"Zapatas": "Zapatas", "Vigas de conexion": "Vigas de conexión", "Cimiento de escalera": "Cimiento de escalera",
        "Columnas": "Columnas", "Falso piso": "Falso piso", "Vigas": "Vigas", "Losa": "Losa aligerada", "Escalera": "Escalera"}
DIAM = ["1/4", "3/8", "1/2", "5/8", "3/4"]

# ------------------------------------------------------------------ estilos
F = "Arial"
f_n = Font(name=F, size=10); f_b = Font(name=F, size=10, bold=True)
f_h = Font(name=F, size=10, bold=True, color="FFFFFF"); f_t = Font(name=F, size=14, bold=True, color="1F3864")
f_st = Font(name=F, size=11, bold=True, color="1F3864"); f_in = Font(name=F, size=10, color="0000FF")
f_lnk = Font(name=F, size=10, color="008000"); f_note = Font(name=F, size=9, italic=True, color="595959")
FILL_H = PatternFill("solid", fgColor="1F3864"); FILL_Y = PatternFill("solid", fgColor="FFFF00")
FILL_T = PatternFill("solid", fgColor="D9E1F2"); FILL_TT = PatternFill("solid", fgColor="B4C6E7")
FILL_L = PatternFill("solid", fgColor="F2F2F2"); FILL_G = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF"); BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)
NUM2 = '#,##0.00;(#,##0.00);"-"'; NUM1 = '#,##0.0;(#,##0.0);"-"'; NUM0 = '#,##0;(#,##0);"-"'
SOL = '"S/ "#,##0.00;("S/ "#,##0.00);"-"'; PCT = '0.0%;(0.0%);"-"'; DATE = "dd/mm/yyyy"

wb = Workbook()
def sheet(name, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws.sheet_view.showGridLines = False
    return ws
def put(ws, r, c, v, font=f_n, fill=None, fmt=None, al=None, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fill: cell.fill = fill
    if fmt: cell.number_format = fmt
    if al: cell.alignment = al
    if border: cell.border = BRD
    return cell
def header(ws, r, c0, labels, height=30):
    for i, t in enumerate(labels):
        put(ws, r, c0 + i, t, f_h, FILL_H, al=CEN)
    ws.row_dimensions[r].height = height
def title(ws, t, sub):
    put(ws, 1, 1, t, f_t, border=False); put(ws, 2, 1, sub, f_note, border=False)
def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[CL(i)].width = x
def inp(ws, r, c, v, fmt=None):
    return put(ws, r, c, v, f_in, FILL_Y, fmt)

# ================================================================== 0. hojas (orden de pestañas)
wsP = sheet("PARÁMETROS", True)
wsM = sheet("METRADO POR NIVEL")
wsA = sheet("ACERO – DESPIECE")
wsC = sheet("LISTA DE COMPRAS")
wsK = sheet("CRONOGRAMA")
wsR = sheet("AVANCE REAL")
wsD = sheet("DASHBOARD")
wsN = sheet("NOTAS")
wsPC = sheet("PLAN DE CORTE")
wsDM = sheet("DATOS MODELO")
SP, SM, SA, SC, SK, SR, SD, SPC, SDM = ("'PARÁMETROS'", "'METRADO POR NIVEL'", "'ACERO – DESPIECE'", "'LISTA DE COMPRAS'",
                                         "CRONOGRAMA", "'AVANCE REAL'", "DASHBOARD", "'PLAN DE CORTE'", "'DATOS MODELO'")

# ================================================================== 1. DATOS MODELO (solo lectura del IFC)
title(wsDM, "DATOS DEL MODELO (lectura del IFC, sin modificar)",
      "Fuente: 'Proyect sj.ifc' exportado de Revit 2027 (27-09-2026). Valores medidos en la geometría; no editar.")
hdr = ["ID Revit", "Nivel", "Elemento", "Tipo Revit", "Comentarios (marcador SJ_*)", "Volumen neto (m³)",
       "Encofrado costados (m²)", "Encofrado fondo (m²)", "Sup. superior libre (m²)", "Costados libres (m²)",
       "Área losa (m²)", "Área aligerada (m²)"]
header(wsDM, 4, 1, hdr)
losas = {v["nivel"]: v for v in AR["losas"].values()}
r = 5
order = sorted(Q, key=lambda x: (NIV.index(nv(x["nivel"])), x["partida"], x["id"]))
for x in order:
    n = nv(x["nivel"]); el = ELEM[x["partida"]]
    vals = [x["id"], n, el, x["tipo"], x["com"], x["vol"], x["enc_lat"], x["enc_fondo"], x["sup_libre"], x["lat_libre_total"]]
    if el == "Losa aligerada":
        vals += [losas[n]["area_losa"], losas[n]["area_alig"]]
    else:
        vals += [None, None]
    for c, v in enumerate(vals, 1):
        put(wsDM, r, c, v, fmt=NUM2 if c >= 6 else None)
    r += 1
DM_LAST = r - 1
put(wsDM, r, 5, "TOTAL", f_b, FILL_TT)
for c in range(6, 13):
    put(wsDM, r, c, f"=SUM({CL(c)}5:{CL(c)}{DM_LAST})", f_b, FILL_TT, NUM2)
r += 2
# acero agregado por nivel / elemento / diametro
put(wsDM, r, 1, "ACERO DEL MODELO por nivel, elemento y diámetro (longitud real de cada barra: incluye ganchos y dobleces)", f_st, border=False)
r += 1
header(wsDM, r, 1, ["Nivel", "Elemento", "Ø", "N° barras", "Longitud total (m)"])
AE0 = r + 1
r += 1
for x in AE:
    for c, v in enumerate([nv(x["nivel"]), ELEM[x["elemento"]], x["d"], x["n"], x["m"]], 1):
        put(wsDM, r, c, v, fmt=NUM2 if c == 5 else None)
    r += 1
AE1 = r - 1
put(wsDM, r, 3, "TOTAL", f_b, FILL_TT); put(wsDM, r, 4, f"=SUM(D{AE0}:D{AE1})", f_b, FILL_TT, NUM0)
put(wsDM, r, 5, f"=SUM(E{AE0}:E{AE1})", f_b, FILL_TT, NUM2)
r += 2
put(wsDM, r, 1, "ÁREAS PARA SOLADO (huella en planta, m²)", f_st, border=False)
put(wsDM, r + 1, 1, "Bajo zapatas"); put(wsDM, r + 1, 2, AR["solado_zap"], fmt=NUM2)
put(wsDM, r + 2, 1, "Bajo vigas de conexión"); put(wsDM, r + 2, 2, AR["solado_vc"], fmt=NUM2)
SOLADO_ZAP = f"{SDM}!$B${r + 1}"; SOLADO_VC = f"{SDM}!$B${r + 2}"
widths(wsDM, [10, 14, 20, 34, 52, 12, 12, 12, 12, 12, 11, 11])
wsDM.freeze_panes = "A5"
DMR = lambda col: f"{SDM}!${col}$5:${col}${DM_LAST}"
AER = lambda col: f"{SDM}!${col}${AE0}:${col}${AE1}"

# ================================================================== 2. PARÁMETROS
title(wsP, "PARÁMETROS DEL CÁLCULO Y PRECIOS UNITARIOS (CUSCO)",
      "Celdas AMARILLAS con texto azul = datos editables. Todo el libro se recalcula al cambiarlas.")
P = {}
r = 4
put(wsP, r, 1, "1. DATOS GENERALES", f_st, border=False); r += 1
header(wsP, r, 1, ["Parámetro", "Valor", "Unidad", "Fuente / observación"]); r += 1
gen = [
    ("fecha_inicio", "Fecha de inicio de obra (casco)", datetime.date(2026, 10, 5), "fecha", "Dato de obra: editar", DATE),
    ("fecha_corte", "Fecha de corte del avance", datetime.date(2026, 10, 5), "fecha", "Actualizar cada semana (DASHBOARD)", DATE),
    ("desp_conc", "Desperdicio de concreto", 0.05, "%", "Dato del usuario (5 %)", PCT),
    ("desp_cas", "Desperdicio / rotura de casetones", 0.03, "%", "Supuesto: manipuleo de EPS", PCT),
    ("usos_mad", "N° de usos de la madera", 4, "usos", "Editable (práctica usual 3-5)", NUM0),
    ("usos_tri", "N° de usos del triplay", 5, "usos", "Editable (triplay 18 mm)", NUM0),
    ("usos_pun", "N° de usos de los puntales", 8, "usos", "Editable (eucalipto rollizo)", NUM0),
    ("alam16", "Alambre N°16 por kg de acero", 0.030, "kg/kg", "APU 'Acero corrugado fy=4200' (CAPECO): 0.030 kg/kg", '0.000'),
    ("dados_zap", "Dados por zapata", 9, "und/zapata", "Supuesto: malla 3x3", NUM0),
    ("cortes_disco", "Cortes por disco de 7\"", 150, "cortes", "Estimado; 1 corte por pieza", NUM0),
    ("curador_lm2", "Consumo de curador químico", 0.20, "L/m²", "Rendimiento típico de membranas de curado (≈5 m²/L)", NUM2),
    ("desmol_lm2", "Consumo de desmoldante", 0.08, "L/m²", "Estimado (≈1 gal por 50 m²)", NUM2),
    ("solado_inc", "Incluir solado (1 = sí, 0 = no)", 1, "1/0", "Partida evitable: no está en el modelo", NUM0),
    ("solado_e", "Espesor del solado", 0.10, "m", "Supuesto", NUM2),
    ("stock_umbral", "Umbral de stock bajo", 0.15, "% de lo pendiente", "Alerta en DASHBOARD", PCT),
    ("lag_enc", "Desfase inicio encofrado tras inicio de acero", 2, "días", "Cronograma", NUM0),
    ("tr34", "Traslape Ø3/4\" (barras > 9 m)", 1.15, "m", "Lámina E-10, nota 8 (sup. clase B). Informativo: el plan de corte ya lo usa", NUM2),
]
for k, lab, v, u, src, fmt in gen:
    put(wsP, r, 1, lab); inp(wsP, r, 2, v, fmt); put(wsP, r, 3, u); put(wsP, r, 4, src, f_note)
    P[k] = f"{SP}!$B${r}"; r += 1
r += 1
put(wsP, r, 1, "2. LOSA ALIGERADA (casetón de poliestireno EPS; el modelo trae losa maciza de 200 mm por criterio)", f_st, border=False); r += 1
header(wsP, r, 1, ["Parámetro", "Valor", "Unidad", "Fuente / observación"]); r += 1
alig = [("h_losa", "Altura total de losa", 0.20, "m", "Láminas E-04 a E-07"),
        ("e_losita", "Espesor de losita", 0.05, "m", "Láminas E-04 a E-07"),
        ("b_vig", "Ancho de vigueta", 0.10, "m", "Láminas E-04 a E-07"),
        ("a_cas", "Ancho del casetón", 0.30, "m", "Casetón EPS 1.20 x 0.30 x 0.15"),
        ("h_cas", "Alto del casetón", 0.15, "m", "Casetón EPS 1.20 x 0.30 x 0.15"),
        ("l_cas", "Largo del casetón", 1.20, "m", "Casetón EPS 1.20 x 0.30 x 0.15")]
for k, lab, v, u, src in alig:
    put(wsP, r, 1, lab); inp(wsP, r, 2, v, NUM2); put(wsP, r, 3, u); put(wsP, r, 4, src, f_note)
    P[k] = f"{SP}!$B${r}"; r += 1
put(wsP, r, 1, "Concreto por m² de aligerado", f_b)
put(wsP, r, 2, f"={P['e_losita']}+{P['b_vig']}*{P['h_cas']}/({P['b_vig']}+{P['a_cas']})", f_b, fmt='0.0000')
put(wsP, r, 3, "m³/m²"); put(wsP, r, 4, "losita + vigueta x alto casetón / separación de viguetas", f_note)
P["c_alig"] = f"{SP}!$B${r}"; r += 1
put(wsP, r, 1, "Casetones por m² de aligerado", f_b)
put(wsP, r, 2, f"=1/(({P['b_vig']}+{P['a_cas']})*{P['l_cas']})", f_b, fmt='0.000')
put(wsP, r, 3, "und/m²"); put(wsP, r, 4, "1 / (separación de viguetas x largo del casetón)", f_note)
P["n_cas"] = f"{SP}!$B${r}"; r += 2

put(wsP, r, 1, "3. f'c POR ELEMENTO (no existe en el modelo: dato de las especificaciones, láminas E-01 a E-10)", f_st, border=False); r += 1
header(wsP, r, 1, ["Elemento", "f'c (kg/cm²)", "", "Fuente / observación"]); r += 1
FC0 = r
for el, fc in [("Zapatas", 210), ("Vigas de conexión", 210), ("Cimiento de escalera", 210), ("Columnas", 210),
               ("Vigas", 210), ("Losa aligerada", 210), ("Escalera", 210), ("Falso piso", 175), ("Solado", 100)]:
    put(wsP, r, 1, el); inp(wsP, r, 2, fc, NUM0); put(wsP, r, 3, ""); put(wsP, r, 4, "Especificación del proyecto (nota E-10)", f_note); r += 1
FC1 = r - 1
FCR = (f"{SP}!$A${FC0}:$A${FC1}", f"{SP}!$B${FC0}:$B${FC1}")
r += 1
put(wsP, r, 1, "4. DOSIFICACIÓN POR m³ DE CONCRETO (cemento en bolsas de 42.5 kg, piedra chancada 1/2\")", f_st, border=False); r += 1
header(wsP, r, 1, ["f'c (kg/cm²)", "Cemento (bls/m³)", "Arena gruesa (m³/m³)", "Piedra chancada (m³/m³)", "Agua (L/m³)", "Fuente"]); r += 1
DO0 = r
for fc, b, a, p_, w, src in [(100, 5.20, 0.56, 0.68, 182, "Estimado (solado): verificar"),
                             (140, 7.01, 0.51, 0.64, 184, "Tabla CAPECO, 'Costos y Presupuestos en Edificación'"),
                             (175, 8.43, 0.54, 0.55, 185, "Tabla CAPECO"), (210, 9.73, 0.52, 0.53, 186, "Tabla CAPECO"),
                             (245, 11.50, 0.50, 0.51, 187, "Tabla CAPECO"), (280, 13.34, 0.45, 0.51, 189, "Tabla CAPECO")]:
    put(wsP, r, 1, fc, f_b, fmt=NUM0)
    for c, v in zip(range(2, 6), [b, a, p_, w]):
        inp(wsP, r, c, v, NUM2)
    put(wsP, r, 6, src, f_note); r += 1
DO1 = r - 1
DOS = lambda col: f"{SP}!${col}${DO0}:${col}${DO1}"
DOSK = f"{SP}!$A${DO0}:$A${DO1}"
r += 1
put(wsP, r, 1, "5. RATIOS DE ENCOFRADO por m² de área de contacto (madera, triplay y puntales se dividen entre sus N° de usos)", f_st, border=False); r += 1
header(wsP, r, 1, ["Tipo de encofrado", "Madera tornillo (p²/m²)", "Triplay 18 mm (m²/m²)", "Clavos 3\" (kg/m²)",
                   "Alambre N°8 (kg/m²)", "Puntales (und/m²)", "Dados (und/m²)", "Fuente"]); r += 1
EN0 = r
ENC_T = [("VC costados", 3.5, 0, 0.15, 0.20, 0, 2), ("Cimiento escalera", 3.5, 0, 0.15, 0.20, 0, 2),
         ("Columnas", 3.0, 1.05, 0.15, 0.25, 0, 2), ("Vigas costados", 3.0, 1.05, 0.18, 0.25, 0, 2),
         ("Vigas fondo", 4.0, 1.05, 0.18, 0.25, 3.5, 4), ("Losa fondo", 2.5, 1.05, 0.12, 0.15, 1.0, 2.5),
         ("Losa bordes", 3.0, 0, 0.15, 0.20, 0, 0), ("Escalera fondo", 4.0, 1.05, 0.20, 0.20, 1.2, 2.5),
         ("Escalera costados", 3.5, 0, 0.20, 0.20, 0, 0)]
for t in ENC_T:
    put(wsP, r, 1, t[0])
    for c, v in enumerate(t[1:], 2):
        inp(wsP, r, c, v, NUM2)
    put(wsP, r, 8, "Referencial de APU usuales (CAPECO); ajustar con el maestro de obra", f_note); r += 1
EN1 = r - 1
ENK = f"{SP}!$A${EN0}:$A${EN1}"
ENC = lambda col: f"{SP}!${col}${EN0}:${col}${EN1}"
r += 1
put(wsP, r, 1, "6. PRECIOS UNITARIOS REFERENCIALES - CUSCO (incluyen IGV; confirmar precio puesto en obra, San Jerónimo)", f_st, border=False); r += 1
header(wsP, r, 1, ["Insumo", "Unidad de compra", "Precio (S/)", "Fecha de cotización", "Proveedor / fuente", "Categoría", "Observación"]); r += 1
PR0 = r
F1 = datetime.date(2026, 9, 28)
PRECIOS = [
    ("Cemento Portland IP 42.5 kg", "bolsa", 27.00, "Sodimac.pe - Cemento Yura Tipo IP", "Concreto", "En Cusco hay distribuidores desde ~S/ 23 (almacén); confirmar"),
    ("Arena gruesa", "m³", 70.00, "Estimado: cotización 'Agregados Cusco' S/ 60 + rango 2026 S/ 55-70 (OneEstimate)", "Concreto", "Estimado: cotizar"),
    ("Piedra chancada 1/2\"", "m³", 80.00, "Estimado: 'Agregados Cusco' S/ 65 + rango 2026 S/ 75-95 (OneEstimate)", "Concreto", "Estimado: cotizar"),
    ("Agua", "m³", 7.00, "Estimado (tarifa no doméstica SEDACUSCO)", "Concreto", "Estimado: confirmar"),
    ("Casetón EPS 1.20x0.30x0.15", "und", 12.90, "Promart.pe - Etsapol (tienda Cusco)", "Concreto", "Ref. Dinova S/ 12.00"),
    ("Varilla Ø1/4\" x 9 m", "und", 9.50, "Estimado (2.25 kg x ~S/ 4.2/kg)", "Acero", "Estimado: cotizar (Aceros Arequipa)"),
    ("Varilla Ø3/8\" x 9 m", "und", 20.13, "Promart.pe - Aceros Arequipa A615, precio online", "Acero", ""),
    ("Varilla Ø1/2\" x 9 m", "und", 36.09, "Promart.pe - Aceros Arequipa A615, precio online", "Acero", ""),
    ("Varilla Ø5/8\" x 9 m", "und", 55.82, "Promart.pe - Aceros Arequipa A615, precio online", "Acero", ""),
    ("Varilla Ø3/4\" x 9 m", "und", 82.09, "Promart.pe - Aceros Arequipa A615, precio online", "Acero", ""),
    ("Alambre negro N°16", "kg", 3.65, "Promart.pe - Prodac rollo 100 kg (S/ 365)", "Acero", ""),
    ("Disco de corte 7\"", "und", 12.00, "Estimado", "Acero", "Estimado: cotizar"),
    ("Dados de concreto", "und", 0.50, "Estimado (fabricados en obra)", "Acero", "Estimado"),
    ("Madera tornillo", "p²", 5.50, "Estimado: cotización madera tornillo S/ 4.90-4.95/p² (antigua) actualizada", "Encofrado", "Estimado: cotizar en madereras de San Sebastián"),
    ("Triplay 18 mm 1.22x2.44", "plancha", 115.00, "Estimado", "Encofrado", "Estimado: cotizar"),
    ("Puntal eucalipto 3\"-4\" x 3 m", "und", 18.00, "Estimado (ref. rollizo 5\" x 3 m S/ 49, IMA)", "Encofrado", "Estimado: cotizar"),
    ("Clavos 3\"", "kg", 6.50, "Estimado", "Encofrado", "Estimado: cotizar"),
    ("Alambre negro N°8", "kg", 5.90, "Promart.pe - Prodac rollo 10 kg (S/ 59)", "Encofrado", ""),
    ("Desmoldante", "gal", 45.00, "Estimado", "Encofrado", "Estimado: cotizar"),
    ("Curador químico (membrana)", "gal", 40.00, "Estimado", "Curado", "Estimado: cotizar"),
]
for ins, u, pu, src, cat, obs in PRECIOS:
    put(wsP, r, 1, ins); put(wsP, r, 2, u); inp(wsP, r, 3, pu, SOL); inp(wsP, r, 4, F1, DATE)
    inp(wsP, r, 5, src); put(wsP, r, 6, cat); put(wsP, r, 7, obs, f_note); r += 1
PR1 = r - 1
PRK = f"{SP}!$A${PR0}:$A${PR1}"
PRC = lambda col: f"{SP}!${col}${PR0}:${col}${PR1}"
r += 1
put(wsP, r, 1, "7. PESO NOMINAL DEL ACERO (NTP 341.031 / ASTM A615)", f_st, border=False); r += 1
header(wsP, r, 1, ["Ø", "kg/m", "Insumo de compra", ""]); r += 1
KG0 = r
for d, k in [("1/4", 0.250), ("3/8", 0.560), ("1/2", 0.994), ("5/8", 1.552), ("3/4", 2.235)]:
    put(wsP, r, 1, d, f_b); inp(wsP, r, 2, k, '0.000'); put(wsP, r, 3, f'Varilla Ø{d}" x 9 m'); r += 1
KG1 = r - 1
KGK = f"{SP}!$A${KG0}:$A${KG1}"; KGV = f"{SP}!$B${KG0}:$B${KG1}"
def kgm(d): return f"INDEX({KGV},MATCH(\"{d}\",{KGK},0))"
r += 1
put(wsP, r, 1, "Leyenda: amarillo + azul = dato editable · negro = fórmula · verde = vínculo a otra hoja.", f_note, border=False)
widths(wsP, [38, 16, 18, 20, 44, 18, 44, 40])
wsP.freeze_panes = "A4"

# ================================================================== 3. METRADO POR NIVEL
title(wsM, "METRADO POR NIVEL - CASCO ESTRUCTURAL",
      "Cantidades del modelo IFC (hoja DATOS MODELO) + parámetros. Columnas I a W: insumos que genera cada partida.")
MH = ["Nivel", "Partida", "Elemento", "Unidad", "Cantidad neta", "Desperdicio", "Cantidad total", "f'c / tipo",
      "Cemento (bls)", "Arena gruesa (m³)", "Piedra 1/2\" (m³)", "Agua (m³)", "Casetón EPS (und)", "Alambre N°16 (kg)",
      "Madera tornillo (p²)", "Triplay (m²)", "Clavos 3\" (kg)", "Alambre N°8 (kg)", "Puntales (und)", "Desmoldante (gal)",
      "Curador (gal)", "Dados (und)", "Discos de corte (und)", "Nota"]
MC = {h: i + 1 for i, h in enumerate(MH)}
r = 4
def dm_sum(col, niv, el):
    return f"SUMIFS({DMR(col)},{DMR('B')},\"{niv}\",{DMR('C')},\"{el}\")"
def conc_row(r, niv, el, expr, nota=""):
    put(wsM, r, 1, niv); put(wsM, r, 2, "Concreto"); put(wsM, r, 3, el); put(wsM, r, 4, "m³")
    put(wsM, r, 5, "=" + expr, fmt=NUM2); put(wsM, r, 6, f"={P['desp_conc']}", f_lnk, fmt=PCT)
    put(wsM, r, 7, f"=E{r}*(1+F{r})", fmt=NUM2)
    put(wsM, r, 8, f"=INDEX({FCR[1]},MATCH(C{r},{FCR[0]},0))", f_lnk, fmt=NUM0)
    for c, col in zip(range(9, 13), "BCDE"):
        f = f"=G{r}*INDEX({DOS(col)},MATCH(H{r},{DOSK},0))"
        if col == "E": f += "/1000"
        put(wsM, r, c, f, fmt=NUM2)
    put(wsM, r, 24, nota, f_note)
def enc_row(r, niv, el, tipo, expr, nota=""):
    put(wsM, r, 1, niv); put(wsM, r, 2, "Encofrado"); put(wsM, r, 3, el); put(wsM, r, 4, "m²")
    put(wsM, r, 5, "=" + expr, fmt=NUM2); put(wsM, r, 6, 0, fmt=PCT); put(wsM, r, 7, f"=E{r}*(1+F{r})", fmt=NUM2)
    put(wsM, r, 8, tipo)
    L = lambda col: f"INDEX({ENC(col)},MATCH(H{r},{ENK},0))"
    put(wsM, r, 15, f"=G{r}*{L('B')}/{P['usos_mad']}", fmt=NUM1)
    put(wsM, r, 16, f"=G{r}*{L('C')}/{P['usos_tri']}", fmt=NUM2)
    put(wsM, r, 17, f"=G{r}*{L('D')}", fmt=NUM2)
    put(wsM, r, 18, f"=G{r}*{L('E')}", fmt=NUM2)
    put(wsM, r, 19, f"=G{r}*{L('F')}/{P['usos_pun']}", fmt=NUM1)
    put(wsM, r, 20, f"=G{r}*{P['desmol_lm2']}/3.785", fmt=NUM2)
    put(wsM, r, 22, f"=G{r}*{L('G')}", fmt=NUM0)
    put(wsM, r, 24, nota, f_note)
def acero_row(r, niv, el, nota=""):
    put(wsM, r, 1, niv); put(wsM, r, 2, "Acero"); put(wsM, r, 3, el); put(wsM, r, 4, "kg")
    terms = "+".join(f"SUMIFS({AER('E')},{AER('A')},\"{niv}\",{AER('B')},\"{el}\",{AER('C')},\"{d}\")*{kgm(d)}" for d in DIAM)
    put(wsM, r, 5, "=" + terms, fmt=NUM1); put(wsM, r, 6, 0, fmt=PCT); put(wsM, r, 7, f"=E{r}*(1+F{r})", fmt=NUM1)
    put(wsM, r, 8, "fy 4200")
    put(wsM, r, 14, f"=G{r}*{P['alam16']}", fmt=NUM1)
    npz = f"SUMIFS({AER('D')},{AER('A')},\"{niv}\",{AER('B')},\"{el}\")"
    put(wsM, r, 23, f"={npz}/{P['cortes_disco']}", fmt=NUM2)
    if el == "Zapatas":
        put(wsM, r, 22, f"=COUNTIFS({DMR('B')},\"{niv}\",{DMR('C')},\"Zapatas\")*{P['dados_zap']}", fmt=NUM0)
    put(wsM, r, 24, nota or "Desperdicio de corte: ver ACERO – DESPIECE (varillas de 9 m)", f_note)
def cur_row(r, niv, el, expr, nota=""):
    put(wsM, r, 1, niv); put(wsM, r, 2, "Curado"); put(wsM, r, 3, el); put(wsM, r, 4, "m²")
    put(wsM, r, 5, "=" + expr, fmt=NUM2); put(wsM, r, 6, 0, fmt=PCT); put(wsM, r, 7, f"=E{r}*(1+F{r})", fmt=NUM2)
    put(wsM, r, 8, "curador")
    put(wsM, r, 21, f"=G{r}*{P['curador_lm2']}/3.785", fmt=NUM2)
    put(wsM, r, 24, nota, f_note)

LEVEL_BLOCK = {}
for niv in NIV:
    put(wsM, r, 1, niv, Font(name=F, size=12, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5597"))
    for c in range(2, len(MH) + 1):
        put(wsM, r, c, None, fill=PatternFill("solid", fgColor="2F5597"))
    r += 1
    header(wsM, r, 1, MH, 42); r += 1
    r0 = r
    if niv == "CIMENTACIÓN":
        els_c = [("Zapatas", "Vaciadas contra el terreno"), ("Vigas de conexión", ""), ("Cimiento de escalera", "")]
        for el, nt in els_c:
            conc_row(r, niv, el, dm_sum("F", niv, el), nt); r += 1
        conc_row(r, niv, "Solado", f"({SOLADO_ZAP}+{SOLADO_VC})*{P['solado_e']}*{P['solado_inc']}",
                 "Partida evitable (no modelada): área bajo zapatas y VC x espesor"); r += 1
        enc_row(r, niv, "Vigas de conexión", "VC costados", dm_sum("G", niv, "Vigas de conexión"), "Fondo sobre solado/terreno"); r += 1
        enc_row(r, niv, "Cimiento de escalera", "Cimiento escalera", dm_sum("G", niv, "Cimiento de escalera")); r += 1
        for el in ["Zapatas", "Vigas de conexión", "Cimiento de escalera"]:
            acero_row(r, niv, el); r += 1
        cur_row(r, niv, "Zapatas", dm_sum("I", niv, "Zapatas"), "Cara superior (costados contra terreno)"); r += 1
        cur_row(r, niv, "Vigas de conexión", dm_sum("I", niv, "Vigas de conexión") + "+" + dm_sum("J", niv, "Vigas de conexión")); r += 1
        cur_row(r, niv, "Cimiento de escalera", dm_sum("I", niv, "Cimiento de escalera") + "+" + dm_sum("J", niv, "Cimiento de escalera")); r += 1
    elif niv == "NIVEL 1":
        conc_row(r, niv, "Columnas", dm_sum("F", niv, "Columnas"), "Columnas de cimentación a N1"); r += 1
        conc_row(r, niv, "Falso piso", dm_sum("F", niv, "Falso piso"), "e = 0.10 (modelo)"); r += 1
        enc_row(r, niv, "Columnas", "Columnas", dm_sum("G", niv, "Columnas")); r += 1
        acero_row(r, niv, "Columnas"); r += 1
        cur_row(r, niv, "Columnas", dm_sum("J", niv, "Columnas")); r += 1
        cur_row(r, niv, "Falso piso", dm_sum("I", niv, "Falso piso")); r += 1
    else:
        conc_row(r, niv, "Columnas", dm_sum("F", niv, "Columnas")); r += 1
        conc_row(r, niv, "Vigas", dm_sum("F", niv, "Vigas"), "Peralte bajo la losa (modelo con uniones)"); r += 1
        al = dm_sum("L", niv, "Losa aligerada"); ls = dm_sum("K", niv, "Losa aligerada")
        conc_row(r, niv, "Losa aligerada", f"{al}*{P['c_alig']}+({ls}-{al})*{P['h_losa']}",
                 "Aligerado x concreto/m² + zonas macizas (sobre vigas y columnas, vigas chatas) x h"); r += 1
        conc_row(r, niv, "Escalera", dm_sum("F", niv, "Escalera"), "Tramo que llega a este nivel"); r += 1
        put(wsM, r, 1, niv); put(wsM, r, 2, "Aligerante"); put(wsM, r, 3, "Losa aligerada"); put(wsM, r, 4, "und")
        put(wsM, r, 5, f"={al}*{P['n_cas']}", fmt=NUM1); put(wsM, r, 6, f"={P['desp_cas']}", f_lnk, fmt=PCT)
        put(wsM, r, 7, f"=E{r}*(1+F{r})", fmt=NUM1); put(wsM, r, 8, "EPS"); put(wsM, r, 13, f"=G{r}", fmt=NUM0)
        put(wsM, r, 24, "Casetones = área aligerada x und/m²", f_note); r += 1
        enc_row(r, niv, "Columnas", "Columnas", dm_sum("G", niv, "Columnas")); r += 1
        enc_row(r, niv, "Vigas", "Vigas costados", dm_sum("G", niv, "Vigas")); r += 1
        enc_row(r, niv, "Vigas", "Vigas fondo", dm_sum("H", niv, "Vigas")); r += 1
        enc_row(r, niv, "Losa aligerada", "Losa fondo", dm_sum("H", niv, "Losa aligerada"), "Fondo libre (sin la huella de vigas y columnas)"); r += 1
        enc_row(r, niv, "Losa aligerada", "Losa bordes", dm_sum("G", niv, "Losa aligerada")); r += 1
        enc_row(r, niv, "Escalera", "Escalera fondo", dm_sum("H", niv, "Escalera")); r += 1
        enc_row(r, niv, "Escalera", "Escalera costados", dm_sum("G", niv, "Escalera"), "Costados y contrapasos"); r += 1
        for el in ["Columnas", "Vigas", "Losa aligerada", "Escalera"]:
            acero_row(r, niv, el); r += 1
        cur_row(r, niv, "Columnas", dm_sum("J", niv, "Columnas")); r += 1
        cur_row(r, niv, "Vigas", dm_sum("J", niv, "Vigas"), "Costados al desencofrar"); r += 1
        cur_row(r, niv, "Losa aligerada", dm_sum("I", niv, "Losa aligerada") + "+" + dm_sum("J", niv, "Losa aligerada")); r += 1
        cur_row(r, niv, "Escalera", dm_sum("I", niv, "Escalera") + "+" + dm_sum("J", niv, "Escalera")); r += 1
    r1 = r - 1
    put(wsM, r, 1, f"TOTAL {niv}", f_b, FILL_TT)
    for c in range(2, 9):
        put(wsM, r, c, None, fill=FILL_TT)
    for c in range(9, 24):
        put(wsM, r, c, f"=SUM({CL(c)}{r0}:{CL(c)}{r1})", f_b, FILL_TT, NUM1)
    put(wsM, r, 24, None, fill=FILL_TT)
    LEVEL_BLOCK[niv] = (r0, r1, r)
    r += 1
    # sub-resumen del nivel
    for par, u in [("Concreto", "m³"), ("Encofrado", "m²"), ("Acero", "kg"), ("Curado", "m²")]:
        put(wsM, r, 2, f"Subtotal {par}", f_b, FILL_T); put(wsM, r, 4, u, fill=FILL_T)
        put(wsM, r, 7, f"=SUMIFS(G{r0}:G{r1},B{r0}:B{r1},\"{par}\")", f_b, FILL_T, NUM2); r += 1
    r += 1
M_LAST = r
widths(wsM, [14, 12, 20, 7, 11, 9, 11, 13] + [10] * 15 + [44])
wsM.freeze_panes = "D4"
MR = lambda col: f"{SM}!${col}$4:${col}${M_LAST}"

# ================================================================== 4. PLAN DE CORTE
title(wsPC, "PLAN DE CORTE DE VARILLAS DE 9 m (optimizado por nivel y diámetro)",
      "Mejor ajuste decreciente sobre la longitud real de cada barra del modelo. Barras > 9 m: 9.00 + (L − 9 + traslape 1.15).")
header(wsPC, 4, 1, ["Nivel", "Ø", "Varilla N°", "Piezas a cortar (m)", "N° piezas", "Longitud usada (m)", "Retazo (m)",
                    "Retazo reutilizable (≥ 1.00 m)", "Elementos", "IDs de barra (Revit)"])
r = 5
plan = sorted(CORTE["plan"], key=lambda p: (NIV.index(nv(p["nivel"])), DIAM.index(p["d"]), p["varilla"]))
for p in plan:
    vals = [nv(p["nivel"]), p["d"], p["varilla"], " + ".join(f"{x:.2f}" for x in p["piezas"]), len(p["piezas"]), p["usado"]]
    for c, v in enumerate(vals, 1):
        cell = put(wsPC, r, c, v, fmt=NUM2 if c == 6 else None)
    put(wsPC, r, 7, f"=9-F{r}", fmt=NUM2)
    put(wsPC, r, 8, f"=IF(G{r}>=1,G{r},0)", fmt=NUM2)
    put(wsPC, r, 9, ", ".join(p["elementos"])); put(wsPC, r, 10, ", ".join(p["ids"]))
    r += 1
PC_LAST = r - 1
widths(wsPC, [14, 6, 9, 46, 9, 12, 10, 14, 30, 60])
wsPC.freeze_panes = "A5"
PCR = lambda col: f"{SPC}!${col}$5:${col}${PC_LAST}"
r += 1
put(wsPC, r, 1, "BARRAS DEL MODELO MAYORES DE 9 m (se cortan en dos piezas con traslape; no se modifica el diseño)", f_st, border=False); r += 1
header(wsPC, r, 1, ["Nivel", "Ø", "ID Revit", "Longitud modelo (m)", "Pieza 1 (m)", "Pieza 2 = L − 9 + traslape (m)"]); r += 1
for x in sorted(CORTE["largas"], key=lambda x: (NIV.index(nv(x["nivel"])), x["id"])):
    for c, v in enumerate([nv(x["nivel"]), x["d"], x["id"], x["L"], x["p1"]], 1):
        put(wsPC, r, c, v, fmt=NUM2 if c >= 4 else None)
    put(wsPC, r, 6, f"=D{r}-9+{P['tr34']}", fmt=NUM2); r += 1

# ================================================================== 5. ACERO – DESPIECE
title(wsA, "ACERO – DESPIECE POR DIÁMETRO Y NIVEL (varillas comerciales de 9 m)",
      "Piezas y longitudes del modelo; varillas y retazos del PLAN DE CORTE. Unidad de compra: VARILLA de 9 m.")
AH = ["Nivel", "Ø", "N° barras (modelo)", "Longitud neta (m)", "kg netos", "Varillas de 9 m (und)", "Longitud comprada (m)",
      "Longitud cortada (m)", "Traslape adicional (m)", "Retazos (m)", "Retazos reutilizables ≥ 1 m (m)", "Desperdicio real (%)",
      "kg comprados (referencia)"]
header(wsA, 4, 1, AH, 42)
r = 5
A_ROWS = []
for niv in NIV:
    ra = r
    for d in DIAM:
        put(wsA, r, 1, niv); put(wsA, r, 2, d, f_b, al=Alignment(horizontal="center"))
        put(wsA, r, 3, f"=SUMIFS({AER('D')},{AER('A')},A{r},{AER('C')},B{r})", f_lnk, fmt=NUM0)
        put(wsA, r, 4, f"=SUMIFS({AER('E')},{AER('A')},A{r},{AER('C')},B{r})", f_lnk, fmt=NUM2)
        put(wsA, r, 5, f"=D{r}*INDEX({KGV},MATCH(B{r},{KGK},0))", fmt=NUM1)
        put(wsA, r, 6, f"=COUNTIFS({PCR('A')},A{r},{PCR('B')},B{r})", f_lnk, fmt=NUM0)
        put(wsA, r, 7, f"=F{r}*9", fmt=NUM2)
        put(wsA, r, 8, f"=SUMIFS({PCR('F')},{PCR('A')},A{r},{PCR('B')},B{r})", f_lnk, fmt=NUM2)
        put(wsA, r, 9, f"=H{r}-D{r}", fmt=NUM2)
        put(wsA, r, 10, f"=G{r}-H{r}", fmt=NUM2)
        put(wsA, r, 11, f"=SUMIFS({PCR('H')},{PCR('A')},A{r},{PCR('B')},B{r})", f_lnk, fmt=NUM2)
        put(wsA, r, 12, f"=IF(G{r}>0,J{r}/G{r},0)", fmt=PCT)
        put(wsA, r, 13, f"=G{r}*INDEX({KGV},MATCH(B{r},{KGK},0))", fmt=NUM1)
        r += 1
    put(wsA, r, 1, f"Subtotal {niv}", f_b, FILL_T); put(wsA, r, 2, None, fill=FILL_T)
    for c in range(3, 14):
        if c == 12:
            put(wsA, r, c, f"=IF(G{r}>0,J{r}/G{r},0)", f_b, FILL_T, PCT)
        else:
            put(wsA, r, c, f"=SUM({CL(c)}{ra}:{CL(c)}{r - 1})", f_b, FILL_T, NUM1 if c in (5, 13) else NUM2 if c > 3 else NUM0)
    A_ROWS.append((ra, r - 1)); r += 1
r += 1
put(wsA, r, 1, "TOTAL POR DIÁMETRO", f_st, border=False); r += 1
header(wsA, r, 1, AH, 42); r += 1
AT0 = r
for d in DIAM:
    put(wsA, r, 1, "TODOS", f_b); put(wsA, r, 2, d, f_b, al=Alignment(horizontal="center"))
    for c in [3, 4, 5, 6, 7, 8, 9, 10, 11, 13]:
        put(wsA, r, c, f"=SUMIFS({CL(c)}$5:{CL(c)}${AT0 - 3},$B$5:$B${AT0 - 3},B{r},$A$5:$A${AT0 - 3},\"<>Subtotal*\")",
            fmt=NUM0 if c in (3, 6) else NUM1 if c in (5, 13) else NUM2)
    put(wsA, r, 12, f"=IF(G{r}>0,J{r}/G{r},0)", fmt=PCT); r += 1
put(wsA, r, 1, "TOTAL GENERAL", f_b, FILL_TT); put(wsA, r, 2, None, fill=FILL_TT)
for c in range(3, 14):
    put(wsA, r, c, f"=IF(G{r}>0,J{r}/G{r},0)" if c == 12 else f"=SUM({CL(c)}{AT0}:{CL(c)}{r - 1})", f_b, FILL_TT,
        PCT if c == 12 else NUM0 if c in (3, 6) else NUM1)
A_TOT = r
r += 2
put(wsA, r, 1, "CONTROL DE CALIDAD contra los conteos del IFC (dato del usuario)", f_st, border=False); r += 1
header(wsA, r, 1, ["Control", "Ø / grupo", "Esperado", "Calculado", "Diferencia", "Estado"]); r += 1
for d, exp in [("3/8", 8847), ("5/8", 948), ("1/4", 912), ("3/4", 520), ("1/2", 372)]:
    put(wsA, r, 1, "N° de barras"); put(wsA, r, 2, d); inp(wsA, r, 3, exp, NUM0)
    put(wsA, r, 4, f"=SUMIFS(C{AT0}:C{AT0 + 4},B{AT0}:B{AT0 + 4},B{r})", fmt=NUM0)
    put(wsA, r, 5, f"=D{r}-C{r}", fmt=NUM0); put(wsA, r, 6, f"=IF(ABS(E{r})<0.5,\"OK\",\"REVISAR\")", f_b); r += 1
for g, niv, exp in [("Vigas de conexión", "CIMENTACIÓN", 580), ("Vigas N2", "NIVEL 2", 721), ("Vigas N3", "NIVEL 3", 561),
                    ("Vigas N4", "NIVEL 4", 549), ("Vigas N5", "NIVEL 5", 538)]:
    el = "Vigas de conexión" if niv == "CIMENTACIÓN" else "Vigas"
    put(wsA, r, 1, "Longitud Ø3/4\" (m)"); put(wsA, r, 2, g); inp(wsA, r, 3, exp, NUM0)
    put(wsA, r, 4, f"=SUMIFS({AER('E')},{AER('A')},\"{niv}\",{AER('B')},\"{el}\",{AER('C')},\"3/4\")", fmt=NUM1)
    put(wsA, r, 5, f"=D{r}-C{r}", fmt=NUM1); put(wsA, r, 6, f"=IF(ABS(E{r})<=1,\"OK\",\"REVISAR\")", f_b); r += 1
put(wsA, r, 1, "Longitud Ø3/4\" total (m)"); put(wsA, r, 2, "Solo en vigas"); inp(wsA, r, 3, 2948, NUM0)
put(wsA, r, 4, f"=SUMIFS({AER('E')},{AER('C')},\"3/4\")", fmt=NUM1); put(wsA, r, 5, f"=D{r}-C{r}", fmt=NUM1)
put(wsA, r, 6, f"=IF(ABS(E{r})<=1,\"OK\",\"REVISAR\")", f_b)
QC1 = r
wsA.conditional_formatting.add(f"F{QC1 - 10}:F{QC1}", CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_G))
wsA.conditional_formatting.add(f"F{QC1 - 10}:F{QC1}", CellIsRule(operator="equal", formula=['"REVISAR"'], fill=PatternFill("solid", fgColor="FFC7CE")))
widths(wsA, [16, 6, 12, 12, 11, 12, 12, 12, 12, 11, 14, 11, 13])
wsA.freeze_panes = "C5"

# ================================================================== 6. LISTA DE COMPRAS
title(wsC, "LISTA DE COMPRAS - CASCO ESTRUCTURAL (unidades comerciales, precios referenciales Cusco)",
      "Acero: unidad de compra = VARILLA de 9 m; los kg van solo como referencia. Alambres en kg.")
INS = [  # insumo, columna METRADO o acero, redondeo
    ("Cemento Portland IP 42.5 kg", ("M", "I"), "und"), ("Arena gruesa", ("M", "J"), "m3"), ("Piedra chancada 1/2\"", ("M", "K"), "m3"),
    ("Agua", ("M", "L"), "m3"), ("Casetón EPS 1.20x0.30x0.15", ("M", "M"), "und"),
    ("Varilla Ø1/4\" x 9 m", ("V", "1/4"), "und"), ("Varilla Ø3/8\" x 9 m", ("V", "3/8"), "und"), ("Varilla Ø1/2\" x 9 m", ("V", "1/2"), "und"),
    ("Varilla Ø5/8\" x 9 m", ("V", "5/8"), "und"), ("Varilla Ø3/4\" x 9 m", ("V", "3/4"), "und"),
    ("Alambre negro N°16", ("M", "N"), "und"), ("Disco de corte 7\"", ("M", "W"), "und"), ("Dados de concreto", ("M", "V"), "und"),
    ("Madera tornillo", ("M", "O"), "und"), ("Triplay 18 mm 1.22x2.44", ("M", "P"), "plancha"), ("Puntal eucalipto 3\"-4\" x 3 m", ("M", "S"), "und"),
    ("Clavos 3\"", ("M", "Q"), "und"), ("Alambre negro N°8", ("M", "R"), "und"), ("Desmoldante", ("M", "T"), "und"),
    ("Curador químico (membrana)", ("M", "U"), "und")]
CH = ["Nivel", "Categoría", "Insumo", "Unidad de compra", "Cantidad requerida", "Cantidad a comprar", "Precio unitario (S/)",
      "Subtotal (S/)", "Peso de referencia (kg)", "Nota"]
r = 4
put(wsC, r, 1, "RESUMEN CONSOLIDADO (todos los niveles)", f_st, border=False); r += 1
header(wsC, r, 1, ["", "Categoría", "Insumo", "Unidad de compra", "Cantidad requerida", "Cantidad a comprar", "Precio unitario (S/)",
                   "Total (S/)", "Peso de referencia (kg)", "% del costo"], 36)
r += 1
CS0 = r
r_cons = {}
for ins, src, rnd in INS:
    r_cons[ins] = r; r += 1
CS1 = r - 1
put(wsC, r, 3, "TOTAL GENERAL (S/)", f_b, FILL_TT)
for c in [1, 2, 4, 5, 6, 7, 9, 10]:
    put(wsC, r, c, None, fill=FILL_TT)
put(wsC, r, 8, f"=SUM(H{CS0}:H{CS1})", Font(name=F, size=11, bold=True), FILL_TT, SOL)
C_TOT = r
r += 2
# subtotales por nivel
put(wsC, r, 1, "SUBTOTAL POR NIVEL Y CATEGORÍA (S/)", f_st, border=False); r += 1
CATS = ["Concreto", "Acero", "Encofrado", "Curado"]
header(wsC, r, 1, ["Nivel"] + CATS + ["Total nivel (S/)", "% del total"]); r += 1
SN0 = r
for niv in NIV:
    put(wsC, r, 1, niv, f_b); r += 1
SN1 = r - 1
put(wsC, r, 1, "TOTAL", f_b, FILL_TT)
for c in range(2, 8):
    put(wsC, r, c, f"=SUM({CL(c)}{SN0}:{CL(c)}{SN1})", f_b, FILL_TT, PCT if c == 7 else SOL)
SN_TOT = r
r += 2
BLK0 = r
blocks = {}
for niv in NIV:
    put(wsC, r, 1, niv, Font(name=F, size=12, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5597"))
    for c in range(2, 11):
        put(wsC, r, c, None, fill=PatternFill("solid", fgColor="2F5597"))
    r += 1
    header(wsC, r, 1, CH, 36); r += 1
    rb0 = r
    for ins, (kind, key), rnd in INS:
        put(wsC, r, 1, niv); put(wsC, r, 2, f"=INDEX({PRC('F')},MATCH(C{r},{PRK},0))", f_lnk); put(wsC, r, 3, ins)
        put(wsC, r, 4, f"=INDEX({PRC('B')},MATCH(C{r},{PRK},0))", f_lnk)
        if kind == "M":
            q = f"=SUMIFS({MR(key)},{MR('A')},A{r})"
            if key == "P":
                q += "/(1.22*2.44)"
            put(wsC, r, 5, q, f_lnk, fmt=NUM2)
        else:
            put(wsC, r, 5, f"=COUNTIFS({PCR('A')},A{r},{PCR('B')},\"{key}\")", f_lnk, fmt=NUM0)
        if rnd == "m3":
            put(wsC, r, 6, f"=CEILING(E{r},0.5)", fmt=NUM1)
        else:
            put(wsC, r, 6, f"=ROUNDUP(E{r},0)", fmt=NUM0)
        put(wsC, r, 7, f"=INDEX({PRC('C')},MATCH(C{r},{PRK},0))", f_lnk, fmt=SOL)
        put(wsC, r, 8, f"=F{r}*G{r}", fmt=SOL)
        if kind == "V":
            put(wsC, r, 9, f"=F{r}*9*{kgm(key)}", fmt=NUM1); put(wsC, r, 10, "kg solo de referencia", f_note)
        elif "Alambre" in ins:
            put(wsC, r, 9, f"=F{r}", fmt=NUM1); put(wsC, r, 10, "Se compra en kg", f_note)
        else:
            put(wsC, r, 9, None); put(wsC, r, 10, "Planchas de 1.22 x 2.44" if key == "P" else
                                         "Madera/triplay/puntales: requerido ÷ N° de usos" if key in "OS" else "", f_note)
        r += 1
    rb1 = r - 1
    put(wsC, r, 3, f"SUBTOTAL {niv}", f_b, FILL_TT)
    for c in [1, 2, 4, 5, 6, 7, 9, 10]:
        put(wsC, r, c, None, fill=FILL_TT)
    put(wsC, r, 8, f"=SUM(H{rb0}:H{rb1})", f_b, FILL_TT, SOL)
    blocks[niv] = (rb0, rb1, r)
    r += 2
C_LAST = r
CR = lambda col: f"${col}${BLK0}:${col}${C_LAST}"
# consolidado
for ins, (kind, key), rnd in INS:
    rr = r_cons[ins]
    put(wsC, rr, 2, f"=INDEX({PRC('F')},MATCH(C{rr},{PRK},0))", f_lnk); put(wsC, rr, 3, ins, f_b)
    put(wsC, rr, 4, f"=INDEX({PRC('B')},MATCH(C{rr},{PRK},0))", f_lnk)
    put(wsC, rr, 5, f"=SUMIFS({CR('E')},{CR('C')},C{rr})", fmt=NUM2)
    put(wsC, rr, 6, f"=SUMIFS({CR('F')},{CR('C')},C{rr})", f_b, fmt=NUM1 if rnd == "m3" else NUM0)
    put(wsC, rr, 7, f"=INDEX({PRC('C')},MATCH(C{rr},{PRK},0))", f_lnk, fmt=SOL)
    put(wsC, rr, 8, f"=SUMIFS({CR('H')},{CR('C')},C{rr})", f_b, fmt=SOL)
    put(wsC, rr, 9, f"=IF(SUMIFS({CR('I')},{CR('C')},C{rr})=0,\"\",SUMIFS({CR('I')},{CR('C')},C{rr}))", fmt=NUM1)
    put(wsC, rr, 10, f"=IF($H${C_TOT}>0,H{rr}/$H${C_TOT},0)", fmt=PCT)
for i, niv in enumerate(NIV):
    rr = SN0 + i
    for j, cat in enumerate(CATS):
        put(wsC, rr, 2 + j, f"=SUMIFS({CR('H')},{CR('A')},$A{rr},{CR('B')},\"{cat}\")", fmt=SOL)
    put(wsC, rr, 6, f"=SUM(B{rr}:E{rr})", f_b, fmt=SOL)
    put(wsC, rr, 7, f"=IF($F${SN_TOT}>0,F{rr}/$F${SN_TOT},0)", fmt=PCT)
widths(wsC, [14, 12, 30, 12, 13, 13, 13, 15, 14, 34])
wsC.freeze_panes = "D4"

# ================================================================== 7. AVANCE REAL (se llena en obra)
title(wsR, "AVANCE REAL - REGISTRO SEMANAL (única hoja que se llena en obra)",
      "Llenar solo las celdas amarillas. Tabla A: % ejecutado ACUMULADO de cada partida. Tabla B: materiales recibidos y usados.")
put(wsR, 3, 1, "Ejemplo de llenado (no se suma): 12/10/2026 | CIMENTACIÓN | Acero | 40 %   ·   12/10/2026 | Cemento Portland IP 42.5 kg | recibido 200 | usado 150",
    f_note, FILL_L, border=False)
header(wsR, 5, 1, ["Fecha", "Semana N°", "Nivel", "Partida", "% ejecutado acumulado", "Observaciones"], 36)
header(wsR, 5, 8, ["Fecha", "Insumo", "Unidad", "Cantidad recibida", "Cantidad usada", "Guía / proveedor", "Observaciones"], 36)
NR = 400
dv_niv = DataValidation(type="list", formula1='"' + ",".join(NIV) + '"', allow_blank=True)
dv_par = DataValidation(type="list", formula1='"Acero,Encofrado,Concreto,Curado"', allow_blank=True)
dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
dv_ins = DataValidation(type="list", formula1=f"={SP}!$A${PR0}:$A${PR1}", allow_blank=True)
for dv in (dv_niv, dv_par, dv_pct, dv_ins):
    wsR.add_data_validation(dv)
for i in range(NR):
    rr = 6 + i
    inp(wsR, rr, 1, None, DATE); put(wsR, rr, 2, f"=IF(A{rr}=\"\",\"\",INT((A{rr}-{P['fecha_inicio']})/7)+1)", fmt=NUM0)
    inp(wsR, rr, 3, None); inp(wsR, rr, 4, None); inp(wsR, rr, 5, None, PCT); inp(wsR, rr, 6, None)
    inp(wsR, rr, 8, None, DATE); inp(wsR, rr, 9, None)
    put(wsR, rr, 10, f"=IF(I{rr}=\"\",\"\",INDEX({PRC('B')},MATCH(I{rr},{PRK},0)))")
    inp(wsR, rr, 11, None, NUM2); inp(wsR, rr, 12, None, NUM2); inp(wsR, rr, 13, None); inp(wsR, rr, 14, None)
dv_niv.add(f"C6:C{5 + NR}"); dv_par.add(f"D6:D{5 + NR}"); dv_pct.add(f"E6:E{5 + NR}"); dv_ins.add(f"I6:I{5 + NR}")
RA = lambda col: f"{SR}!${col}$6:${col}${5 + NR}"
widths(wsR, [12, 9, 15, 12, 13, 30, 3, 12, 32, 9, 12, 12, 20, 26])
wsR.freeze_panes = "A6"

# ================================================================== 8. CRONOGRAMA
title(wsK, "CRONOGRAMA PROGRAMADO - PARTIDAS POR NIVEL (peso % por costo)",
      "Duraciones editables (días calendario). Secuencia: acero → encofrado (desfase) → concreto → curado; el siguiente nivel inicia al terminar el concreto.")
KH = ["N°", "Nivel", "Partida", "Costo (S/)", "Peso (%)", "Duración (días)", "Inicio", "Fin",
      "% programado a la fecha de corte", "% ejecutado a la fecha de corte"]
header(wsK, 5, 1, KH, 42)
NW = 26
put(wsK, 3, 12, "Curva S: fracción programada por semana (fin de semana)", f_st, border=False)
for w in range(NW):
    c = 12 + w
    put(wsK, 4, c, w + 1, f_h, FILL_H, NUM0, CEN)
    put(wsK, 5, c, f"={P['fecha_inicio']}+7*{w + 1}-1", f_h, FILL_H, "dd/mm", CEN)
    wsK.column_dimensions[CL(c)].width = 7
DUR = {"CIMENTACIÓN": [8, 6, 3, 7], "NIVEL 1": [4, 3, 2, 7]}
r = 6
K0 = r
prev_conc = None
for niv in NIV:
    durs = DUR.get(niv, [12, 12, 2, 7])
    base = r
    for j, par in enumerate(["Acero", "Encofrado", "Concreto", "Curado"]):
        put(wsK, r, 1, r - K0 + 1, fmt=NUM0); put(wsK, r, 2, niv); put(wsK, r, 3, par)
        cat = par
        put(wsK, r, 4, f"=SUMIFS({SC}!$B${SN0}:$E${SN1},{SC}!$A${SN0}:$A${SN1},B{r},{SC}!$B${SN0 - 1}:$E${SN0 - 1},C{r})"
            if False else f"=INDEX({SC}!$B${SN0}:$E${SN1},MATCH(B{r},{SC}!$A${SN0}:$A${SN1},0),MATCH(C{r},{SC}!$B${SN0 - 1}:$E${SN0 - 1},0))",
            f_lnk, fmt=SOL)
        inp(wsK, r, 6, durs[j], NUM0)
        if par == "Acero":
            put(wsK, r, 7, f"={P['fecha_inicio']}" if prev_conc is None else f"=H{prev_conc}+1", fmt=DATE)
        elif par == "Encofrado":
            put(wsK, r, 7, f"=G{base}+{P['lag_enc']}", fmt=DATE)
        elif par == "Concreto":
            put(wsK, r, 7, f"=MAX(H{base},H{base + 1})+1", fmt=DATE)
        else:
            put(wsK, r, 7, f"=H{base + 2}+1", fmt=DATE)
        put(wsK, r, 8, f"=G{r}+F{r}-1", fmt=DATE)
        put(wsK, r, 9, f"=MAX(0,MIN(1,({P['fecha_corte']}-G{r}+1)/(H{r}-G{r}+1)))", fmt=PCT)
        put(wsK, r, 10, f"=_xlfn.MAXIFS({RA('E')},{RA('C')},B{r},{RA('D')},C{r},{RA('A')},\"<=\"&{P['fecha_corte']})", fmt=PCT)
        for w in range(NW):
            c = 12 + w
            put(wsK, r, c, f"=MAX(0,MIN(1,({CL(c)}$5-$G{r}+1)/($H{r}-$G{r}+1)))", fmt='0%;;""')
        r += 1
    prev_conc = base + 2
K1 = r - 1
for rr in range(K0, K1 + 1):
    put(wsK, rr, 5, f"=IF($D${K1 + 1}>0,D{rr}/$D${K1 + 1},0)", fmt=PCT)
put(wsK, r, 3, "TOTAL", f_b, FILL_TT); put(wsK, r, 4, f"=SUM(D{K0}:D{K1})", f_b, FILL_TT, SOL)
put(wsK, r, 5, f"=SUM(E{K0}:E{K1})", f_b, FILL_TT, PCT); put(wsK, r, 7, f"=MIN(G{K0}:G{K1})", f_b, FILL_TT, DATE)
put(wsK, r, 8, f"=MAX(H{K0}:H{K1})", f_b, FILL_TT, DATE)
put(wsK, r, 9, f"=SUMPRODUCT($E${K0}:$E${K1},I{K0}:I{K1})", f_b, FILL_TT, PCT)
put(wsK, r, 10, f"=SUMPRODUCT($E${K0}:$E${K1},J{K0}:J{K1})", f_b, FILL_TT, PCT)
K_TOT = r
# grilla de ejecutado (para la curva S)
r += 2
put(wsK, r, 12, "Fracción ejecutada por semana (según AVANCE REAL, hasta la fecha de corte)", f_st, border=False)
r += 1
KE0 = r
for i in range(K1 - K0 + 1):
    rr = r + i; src = K0 + i
    put(wsK, rr, 2, f"=B{src}", f_lnk); put(wsK, rr, 3, f"=C{src}", f_lnk)
    for w in range(NW):
        c = 12 + w
        put(wsK, rr, c, f"=_xlfn.MAXIFS({RA('E')},{RA('C')},$B{rr},{RA('D')},$C{rr},{RA('A')},\"<=\"&MIN({CL(c)}$5,{P['fecha_corte']}))",
            fmt='0%;;""')
KE1 = r + (K1 - K0)
# gantt: sombreado de la grilla programada
wsK.conditional_formatting.add(f"L{K0}:{CL(11 + NW)}{K1}",
                               FormulaRule(formula=[f"AND(L{K0}>0,L{K0}<1)"], fill=PatternFill("solid", fgColor="9BC2E6")))
wsK.conditional_formatting.add(f"L{K0}:{CL(11 + NW)}{K1}",
                               FormulaRule(formula=[f"L{K0}>=1"], fill=PatternFill("solid", fgColor="D9E1F2")))
widths(wsK, [5, 14, 12, 14, 9, 10, 11, 11, 13, 13, 2])
wsK.freeze_panes = "D6"

# ================================================================== 9. DASHBOARD
title(wsD, "DASHBOARD DE CONTROL - PROGRAMADO vs EJECUTADO",
      "Se actualiza solo con la fecha de corte (PARÁMETROS) y el registro de AVANCE REAL. Semáforo por desviación: verde ≤ 5 %, ámbar 5-15 %, rojo > 15 %.")
put(wsD, 4, 1, "Fecha de corte", f_b); put(wsD, 4, 2, f"={P['fecha_corte']}", f_lnk, fmt=DATE)
put(wsD, 5, 1, "Costo total casco (S/)", f_b); put(wsD, 5, 2, f"={SC}!H{C_TOT}", f_lnk, fmt=SOL)
put(wsD, 6, 1, "% programado", f_b); put(wsD, 6, 2, f"={SK}!I{K_TOT}", f_lnk, fmt=PCT)
put(wsD, 7, 1, "% ejecutado", f_b); put(wsD, 7, 2, f"={SK}!J{K_TOT}", f_lnk, fmt=PCT)
put(wsD, 8, 1, "Desviación (prog. − ejec.)", f_b); put(wsD, 8, 2, "=B6-B7", f_b, fmt=PCT)
put(wsD, 9, 1, "Semáforo general", f_b); put(wsD, 9, 2, '=IF(B8<=0.05,"VERDE",IF(B8<=0.15,"ÁMBAR","ROJO"))', f_b, al=CEN)
put(wsD, 10, 1, "Fin programado del casco", f_b); put(wsD, 10, 2, f"={SK}!H{K_TOT}", f_lnk, fmt=DATE)
def semaf(rng):
    wsD.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"VERDE"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, bold=True, color="006100")))
    wsD.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ÁMBAR"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(name=F, bold=True, color="9C5700")))
    wsD.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ROJO"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, bold=True, color="9C0006")))
semaf("B9")
# avance por nivel
r = 12
put(wsD, r, 1, "AVANCE POR NIVEL", f_st, border=False); r += 1
header(wsD, r, 1, ["Nivel", "Peso (%)", "% programado", "% ejecutado", "Desviación", "Semáforo"]); r += 1
DN0 = r
KR = lambda col: f"{SK}!${col}${K0}:${col}${K1}"
for niv in NIV:
    put(wsD, r, 1, niv, f_b)
    put(wsD, r, 2, f"=SUMIFS({KR('E')},{KR('B')},A{r})", fmt=PCT)
    put(wsD, r, 3, f"=IF(B{r}>0,SUMPRODUCT(({KR('B')}=A{r})*{KR('E')}*{KR('I')})/B{r},0)", fmt=PCT)
    put(wsD, r, 4, f"=IF(B{r}>0,SUMPRODUCT(({KR('B')}=A{r})*{KR('E')}*{KR('J')})/B{r},0)", fmt=PCT)
    put(wsD, r, 5, f"=C{r}-D{r}", fmt=PCT)
    put(wsD, r, 6, f'=IF(E{r}<=0.05,"VERDE",IF(E{r}<=0.15,"ÁMBAR","ROJO"))', f_b, al=CEN); r += 1
DN1 = r - 1
semaf(f"F{DN0}:F{DN1}")
r += 1
put(wsD, r, 1, "AVANCE POR PARTIDA", f_st, border=False); r += 1
header(wsD, r, 1, ["Partida", "Peso (%)", "% programado", "% ejecutado", "Desviación", "Semáforo"]); r += 1
DP0 = r
for par in ["Acero", "Encofrado", "Concreto", "Curado"]:
    put(wsD, r, 1, par, f_b)
    put(wsD, r, 2, f"=SUMIFS({KR('E')},{KR('C')},A{r})", fmt=PCT)
    put(wsD, r, 3, f"=IF(B{r}>0,SUMPRODUCT(({KR('C')}=A{r})*{KR('E')}*{KR('I')})/B{r},0)", fmt=PCT)
    put(wsD, r, 4, f"=IF(B{r}>0,SUMPRODUCT(({KR('C')}=A{r})*{KR('E')}*{KR('J')})/B{r},0)", fmt=PCT)
    put(wsD, r, 5, f"=C{r}-D{r}", fmt=PCT)
    put(wsD, r, 6, f'=IF(E{r}<=0.05,"VERDE",IF(E{r}<=0.15,"ÁMBAR","ROJO"))', f_b, al=CEN); r += 1
semaf(f"F{DP0}:F{r - 1}")
r += 1
# curva S
put(wsD, r, 1, "CURVA S (acumulado ponderado por costo)", f_st, border=False); r += 1
header(wsD, r, 1, ["Semana", "Fecha (fin de semana)", "% programado", "% ejecutado"]); r += 1
CV0 = r
for w in range(NW):
    c = CL(12 + w)
    put(wsD, r, 1, w + 1, fmt=NUM0)
    put(wsD, r, 2, f"={SK}!{c}$5", f_lnk, fmt=DATE)
    put(wsD, r, 3, f"=SUMPRODUCT({SK}!$E${K0}:$E${K1},{SK}!{c}${K0}:{c}${K1})", fmt=PCT)
    put(wsD, r, 4, f"=SUMPRODUCT({SK}!$E${K0}:$E${K1},{SK}!{c}${KE0}:{c}${KE1})", fmt=PCT)
    r += 1
CV1 = r - 1
ch = LineChart(); ch.title = "Curva S: programado vs ejecutado"; ch.y_axis.title = "% acumulado"; ch.x_axis.title = "Semana"
ch.y_axis.number_format = "0%"; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.height = 8.5; ch.width = 17
ch.add_data(Reference(wsD, min_col=3, min_row=CV0 - 1, max_row=CV1), titles_from_data=True)
ch.add_data(Reference(wsD, min_col=4, min_row=CV0 - 1, max_row=CV1), titles_from_data=True)
ch.set_categories(Reference(wsD, min_col=1, min_row=CV0, max_row=CV1))
ch.series[0].graphicalProperties.line.solidFill = "1F3864"; ch.series[1].graphicalProperties.line.solidFill = "C00000"
wsD.add_chart(ch, "H4")
bc = BarChart(); bc.type = "col"; bc.title = "% de avance por nivel"; bc.y_axis.number_format = "0%"
bc.y_axis.scaling.min = 0; bc.y_axis.scaling.max = 1; bc.height = 7.5; bc.width = 17
bc.add_data(Reference(wsD, min_col=3, max_col=4, min_row=DN0 - 1, max_row=DN1), titles_from_data=True)
bc.set_categories(Reference(wsD, min_col=1, min_row=DN0, max_row=DN1))
wsD.add_chart(bc, "H22")
# materiales
r += 1
put(wsD, r, 1, "MATERIALES: requerido vs comprado vs consumido (alerta de stock bajo)", f_st, border=False); r += 1
header(wsD, r, 1, ["Insumo", "Unidad", "Requerido (compra)", "Comprado (recibido)", "Consumido (usado)", "Stock en obra",
                   "% comprado", "% consumido", "Alerta"], 36); r += 1
MT0 = r
for ins, _, _ in INS:
    put(wsD, r, 1, ins); put(wsD, r, 2, f"=INDEX({PRC('B')},MATCH(A{r},{PRK},0))", f_lnk)
    put(wsD, r, 3, f"={SC}!F{r_cons[ins]}", f_lnk, fmt=NUM1)
    put(wsD, r, 4, f"=SUMIFS({RA('K')},{RA('I')},A{r})", fmt=NUM1)
    put(wsD, r, 5, f"=SUMIFS({RA('L')},{RA('I')},A{r})", fmt=NUM1)
    put(wsD, r, 6, f"=D{r}-E{r}", fmt=NUM1)
    put(wsD, r, 7, f"=IF(C{r}>0,D{r}/C{r},0)", fmt=PCT); put(wsD, r, 8, f"=IF(C{r}>0,E{r}/C{r},0)", fmt=PCT)
    put(wsD, r, 9, f'=IF(D{r}=0,"SIN RECEPCIÓN",IF(AND(C{r}-E{r}>0,F{r}<{P["stock_umbral"]}*(C{r}-E{r})),"STOCK BAJO","OK"))', f_b, al=CEN)
    r += 1
MT1 = r - 1
wsD.conditional_formatting.add(f"I{MT0}:I{MT1}", CellIsRule(operator="equal", formula=['"STOCK BAJO"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=F, bold=True, color="9C0006")))
wsD.conditional_formatting.add(f"I{MT0}:I{MT1}", CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))
wsD.conditional_formatting.add(f"I{MT0}:I{MT1}", CellIsRule(operator="equal", formula=['"SIN RECEPCIÓN"'], fill=FILL_L))
mc = BarChart(); mc.type = "bar"; mc.title = "Materiales: % comprado y % consumido"; mc.x_axis.number_format = "0%"
mc.height = 11; mc.width = 17
mc.add_data(Reference(wsD, min_col=7, max_col=8, min_row=MT0 - 1, max_row=MT1), titles_from_data=True)
mc.set_categories(Reference(wsD, min_col=1, min_row=MT0, max_row=MT1))
wsD.add_chart(mc, f"K{MT0 - 1}")
widths(wsD, [30, 14, 14, 14, 13, 13, 11, 11, 14])

# ================================================================== 10. NOTAS
title(wsN, "NOTAS, SUPUESTOS Y FUENTES", "Proyecto: Vivienda multifamiliar San Jerónimo, Cusco. Diseño estructural aprobado: el modelo solo se LEYÓ.")
notas = [
    ("FUENTE DE DATOS", ""),
    ("1", "Cantidades leídas del IFC 'Proyect sj.ifc' (Revit 2027, exportado el 27-09-2026), sin transacciones ni cambios en Revit. Elementos identificados por el parámetro Comentarios (SJ_*)."),
    ("2", "Volumen de concreto = volumen neto de la geometría (Revit ya descuenta los cruces: las vigas llegan hasta el fondo de la losa y las columnas hasta el fondo de vigas)."),
    ("3", "Encofrado = área de contacto medida en la geometría: se descuentan las caras que tocan otro elemento de concreto (fondo de losa sobre vigas, caras de columna donde llegan vigas, etc.). Zapatas vaciadas contra el terreno (sin encofrado); fondos de vigas de conexión, cimiento de escalera y falso piso sobre terreno/solado."),
    ("4", "Acero: longitud real de cada barra del modelo (incluye ganchos y dobleces). Control: N° de barras por diámetro y longitudes de Ø3/4\" cuadran exactamente con los conteos del IFC (ver ACERO – DESPIECE). Peso total 18,793 kg (igual a la lámina E-10)."),
    ("5", "Niveles: CIMENTACIÓN = zapatas, vigas de conexión, cimiento de escalera y solado. NIVEL 1 = columnas de cimentación a N1 y falso piso. NIVEL 2 a 5 = columnas que llegan a ese nivel, sus vigas, su losa y el tramo de escalera que llega a él."),
    ("LOSA", ""),
    ("6", "El tipo de losa en Revit es PISO_ESTRUCTURAL_MACIZO_200mm_VOLADO (maciza 0.20) por criterio del modelo; el acero es de aligerado y las láminas E-04 a E-07 la dibujan aligerada. Por decisión del usuario se metra ALIGERADA con casetón de poliestireno (EPS) 1.20 x 0.30 x 0.15: 69.48 m² de aligerado por piso; el resto de la losa (franjas sobre vigas y columnas y las 2 vigas chatas VCH) se metra maciza de 0.20."),
    ("7", "Concreto de aligerado = losita + vigueta x alto de casetón / separación = 0.0875 m³/m² (h = 0.20). Casetones = 1 / (0.40 x 1.20) = 2.083 und/m² + 3 % de rotura."),
    ("CONCRETO", ""),
    ("8", "f'c no existe como parámetro en el modelo (material 'Hormigón, moldeado in situ'): se toma de la especificación del proyecto (210 estructurales, 175 falso piso, 100 solado). Celdas amarillas en PARÁMETROS."),
    ("9", "Dosificación por m³: tabla CAPECO ('Costos y Presupuestos en Edificación'), cemento en bolsas de 42.5 kg. La fila de f'c 100 (solado) es estimada: verificar con diseño de mezcla. Desperdicio 5 %."),
    ("10", "Solado f'c 100 de 0.10: no está modelado; se calcula con la huella de zapatas (44.14 m²) y vigas de conexión (13.36 m²). Partida evitable: se apaga con el interruptor en PARÁMETROS."),
    ("11", "Recubrimientos: el modelo asigna 25 mm en todas las caras de todos los elementos (incluidas zapatas). Se reporta tal cual; no afecta el metrado."),
    ("ACERO", ""),
    ("12", "Corte optimizado por nivel y diámetro en varillas de 9 m (mejor ajuste decreciente). El desperdicio que se muestra es el real del plan de corte (PLAN DE CORTE, una fila por varilla). Los retazos ≥ 1.00 m se pueden reutilizar en dados, ganchos o bastones cortos."),
    ("13", "42 barras superiores de Ø3/4\" en vigas miden 9.10 - 9.32 m (más que la varilla comercial). No se cambia el diseño: se cortan en 9.00 m + (L − 9 + 1.15 m de traslape, lámina E-10 nota 8). La ubicación del empalme la define el ingeniero (fuera de nudos)."),
    ("14", "Ø1/4\": se asume varilla de 9 m como pidió el usuario; en Cusco también se vende en rollo por kg (alambrón)."),
    ("15", "Alambre N°16: 0.030 kg por kg de acero (APU CAPECO 'acero corrugado fy=4200'). Discos de corte: 1 corte por pieza, 150 cortes por disco (estimado). Dados: ratios por m² de encofrado y 9 por zapata (supuestos)."),
    ("ENCOFRADO Y CURADO", ""),
    ("16", "Ratios de madera, triplay, clavos, alambre N°8 y puntales por m² de contacto: valores referenciales de APU usuales (CAPECO), a ajustar con el maestro de obra. Madera, triplay y puntales se dividen entre su N° de usos (editables): lo comprado en un nivel se reutiliza en los siguientes."),
    ("17", "Curado con curador químico (membrana) a 0.20 L/m² sobre las caras expuestas (superior y costados al desencofrar; zapatas solo la cara superior). Alternativa: curado con agua (no incluida)."),
    ("PRECIOS", ""),
    ("18", "Precios referenciales consultados el 28-09-2026 (incluyen IGV). Fierro Aceros Arequipa, alambres Prodac y casetón EPS: precios online de Promart.pe (tienda en Cusco). Cemento Yura IP: Sodimac.pe (S/ 27); distribuidores de Cusco publican desde ~S/ 23 en almacén. Agregados, madera, triplay, puntales, clavos, desmoldante, curador y discos: ESTIMADOS (marcados en la hoja PARÁMETROS). Cotizar puesto en obra antes de comprar."),
    ("CONTROL DE AVANCE", ""),
    ("19", "El % ejecutado se registra ACUMULADO por partida y nivel en AVANCE REAL; el DASHBOARD toma el último valor registrado hasta la fecha de corte. El peso de cada partida es su costo en LISTA DE COMPRAS."),
    ("20", "Semáforo por desviación (programado − ejecutado): verde ≤ 5 %, ámbar 5 - 15 %, rojo > 15 %. Alerta de stock bajo: stock en obra < umbral x (requerido − consumido)."),
    ("21", "El plan de corte y los datos del modelo son valores leídos del IFC; si el modelo cambia, se debe volver a exportar el IFC y regenerar este libro."),
]
r = 4
for k, t in notas:
    if t == "":
        put(wsN, r, 1, k, f_st, border=False); r += 1; continue
    put(wsN, r, 1, k, f_b, al=Alignment(horizontal="center", vertical="top")); put(wsN, r, 2, t, al=WRAP)
    wsN.row_dimensions[r].height = 15 * max(1, len(t) // 110 + 1); r += 1
widths(wsN, [8, 130])

wb.save(OUT)
print("ok", OUT, dict(M_LAST=M_LAST, C_TOT=C_TOT, K_TOT=K_TOT, PC_LAST=PC_LAST))
